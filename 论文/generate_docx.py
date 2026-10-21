# -*- coding: utf-8 -*-
"""
毕业论文 Markdown → Word 排版脚本
符合中国高校本科毕业论文格式规范
"""

import re
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

# ========== 样式配置 ==========
FONT_BODY = "宋体"
FONT_TITLE = "黑体"
FONT_CODE = "Consolas"
FONT_EN = "Times New Roman"

SIZE_BODY = Pt(12)          # 小四
SIZE_H1 = Pt(16)            # 三号
SIZE_H2 = Pt(14)            # 四号
SIZE_H3 = Pt(12)            # 小四
SIZE_CODE = Pt(9)           # 小五
SIZE_ABSTRACT = Pt(12)      # 小四

LINE_SPACING = Pt(22)       # 1.5倍行距 ≈ 22pt
PAGE_TOP = Cm(2.54)
PAGE_BOTTOM = Cm(2.54)
PAGE_LEFT = Cm(3.17)
PAGE_RIGHT = Cm(3.17)

def set_cell_font(cell, text, font_name=FONT_BODY, font_size=Pt(10)):
    """设置单元格字体"""
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = font_size

def create_document():
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = PAGE_TOP
    section.bottom_margin = PAGE_BOTTOM
    section.left_margin = PAGE_LEFT
    section.right_margin = PAGE_RIGHT

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = FONT_BODY
    font.size = SIZE_BODY
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
    pf = style.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    add_page_number(section)
    return doc


def _field(paragraph, instr, placeholder=""):
    """往段落里塞一个 Word 域（TOC / PAGE 等）。"""
    from docx.oxml import OxmlElement
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr_el = OxmlElement('w:instrText')
    instr_el.set(qn('xml:space'), 'preserve')
    instr_el.text = instr
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t')
    txt.text = placeholder
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_sep)
    run._r.append(txt)
    run._r.append(fld_end)
    run.font.name = FONT_BODY
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)


def add_page_number(section):
    """页脚居中页码。"""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_note = p.add_run("— ")
    run_note.font.name = FONT_BODY
    run_note.font.size = Pt(10.5)
    _field(p, "PAGE", "1")
    tail = p.add_run(" —")
    tail.font.name = FONT_BODY
    tail.font.size = Pt(10.5)


def add_toc(doc):
    """插入「目录」页：标题 + TOC 域（打开 Word 后按 F9 更新页码）。"""
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目  录")
    run.font.name = FONT_TITLE
    run.font.size = SIZE_H1
    run.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_TITLE)

    hint = doc.add_paragraph()
    hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hint.add_run("（在 Word 中右键此处 → 更新域，即生成带页码的目录）")
    r.font.name = FONT_BODY
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)

    _field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "右键→更新域，生成目录")
    doc.add_page_break()

def add_cover_page(doc):
    """添加封面页"""
    for _ in range(6):
        doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("本科毕业设计（论文）")
    run.font.name = FONT_TITLE
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_TITLE)
    run.font.size = Pt(26)
    run.bold = True

    doc.add_paragraph("")
    doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("基于ESP32-S3的Modbus通信诊断工具设计与实现")
    run.font.name = FONT_TITLE
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_TITLE)
    run.font.size = Pt(22)
    run.bold = True

    for _ in range(4):
        doc.add_paragraph("")

    info_items = [
        ("学    院：", "____________"),
        ("专    业：", "____________"),
        ("学    号：", "____________"),
        ("姓    名：", "____________"),
        ("指导教师：", "____________"),
    ]
    for label, value in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        run.font.name = FONT_BODY
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
        run.font.size = Pt(14)
        run = p.add_run(value)
        run.font.name = FONT_BODY
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
        run.font.size = Pt(14)

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026年5月")
    run.font.name = FONT_BODY
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
    run.font.size = Pt(14)

    doc.add_page_break()

def add_heading(doc, text, level=1):
    """添加标题"""
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = SIZE_H1
    elif level == 2:
        size = SIZE_H2
    else:
        size = SIZE_H3

    run = p.add_run(text)
    run.font.name = FONT_TITLE
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_TITLE)
    run.font.size = size
    run.bold = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)

    # 手动排版不走 Heading 样式，但目录域要靠「大纲级别」收集条目 ——
    # 不设 outlineLvl 的话，Word 里更新域会得到一个空目录。
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr()
    olvl = OxmlElement('w:outlineLvl')
    olvl.set(qn('w:val'), str(level - 1))
    pPr.append(olvl)
    return p

def add_body_text(doc, text, indent=True):
    """添加正文段落"""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)  # 首行缩进2字符
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
    run.font.size = SIZE_BODY
    return p

def add_code_block(doc, code):
    """添加代码块"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(code)
    run.font.name = FONT_CODE
    run.font.size = SIZE_CODE
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # 添加灰色背景（通过段落底纹）
    shading = p._element.get_or_add_pPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): 'F5F5F5'
    })
    shading.append(shd)
    return p

def add_table(doc, headers, rows):
    """添加表格"""
    table = doc.add_table(rows=len(rows)+1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # 表头
    for i, h in enumerate(headers):
        set_cell_font(table.rows[0].cells[i], h, FONT_BODY, Pt(10))
        table.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 加粗
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True

    # 数据行
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell_font(table.rows[r+1].cells[c], str(val), FONT_BODY, Pt(10))
            table.rows[r+1].cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")  # 表后空行
    return table

def parse_markdown_and_build(doc, md_path):
    """解析Markdown并构建Word文档"""
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_buffer = []
    in_table = False
    table_headers = []
    table_rows = []
    current_text = []

    def flush_text():
        if current_text:
            text = ''.join(current_text).strip()
            if text:
                add_body_text(doc, text)
            current_text.clear()

    def flush_table():
        nonlocal in_table, table_headers, table_rows
        if in_table and table_headers:
            add_table(doc, table_headers, table_rows)
            table_headers = []
            table_rows = []
            in_table = False

    for line in lines:
        line = line.rstrip('\n')

        # 插图：![图 4-1  系统硬件架构](figures/xxx.png)
        m_img = re.match(r'^\s*!\[(.*?)\]\((.+?)\)\s*$', line)
        if m_img and not in_code_block:
            flush_text()
            flush_table()
            caption, rel = m_img.group(1), m_img.group(2)
            img_path = os.path.join(os.path.dirname(os.path.abspath(md_path)), rel)
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(3)
                p.add_run().add_picture(img_path, width=Inches(5.9))
                if caption:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(10)
                    run = cap.add_run(caption)
                    run.font.name = FONT_BODY
                    run.font.size = Pt(10.5)      # 五号
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
            else:
                print(f"  ! 找不到插图 {img_path}，跳过")
            continue

        # 代码块
        if line.startswith('```'):
            if in_code_block:
                flush_text()
                add_code_block(doc, '\n'.join(code_buffer))
                code_buffer = []
                in_code_block = False
            else:
                flush_text()
                flush_table()
                in_code_block = True
            continue

        if in_code_block:
            code_buffer.append(line)
            continue

        # 表格行
        if line.startswith('|') and '|' in line[1:]:
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if all(set(c) <= {'-', ':', ' '} for c in cells):
                # 分隔行，跳过
                continue
            if not in_table:
                flush_text()
                in_table = True
                table_headers = cells
            else:
                table_rows.append(cells)
            continue
        else:
            flush_table()

        # 空行 → 段落分隔
        if not line.strip():
            flush_text()
            continue

        # 目录占位
        if line.strip() == '<!--TOC-->':
            flush_text()
            flush_table()
            add_toc(doc)
            continue

        # 标题
        if line.startswith('# '):
            flush_text()
            title = line[2:].strip()
            # 跳过文档主标题（已在封面）
            if '基于ESP32-S3' in title:
                continue
            if title.startswith('第') and '章' in title:
                doc.add_page_break()          # 每章另起一页
            add_heading(doc, title, 1)
            continue
        if line.startswith('## '):
            flush_text()
            h2 = line[3:].strip()
            if h2 in ('参考文献', '致谢', '附录', '中文摘要', 'Abstract'):
                doc.add_page_break()          # 这几部分各自另起一页，且按一级标题进目录
                add_heading(doc, h2, 1)
            else:
                add_heading(doc, h2, 2)
            continue
        if line.startswith('### '):
            flush_text()
            add_heading(doc, line[4:].strip(), 3)
            continue

        # 分隔线
        if line.strip() == '---':
            flush_text()
            continue

        # 列表项
        if line.startswith('- **') or line.startswith('- '):
            flush_text()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.first_line_indent = Pt(0)
            # 清理markdown格式
            text = line[2:].strip()
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)
            run = p.add_run("• " + text)
            run.font.name = FONT_BODY
            run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
            run.font.size = SIZE_BODY
            continue

        # 引用行
        if line.startswith('>'):
            flush_text()
            text = line[1:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.5)
            run = p.add_run(text)
            run.font.name = FONT_BODY
            run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
            run.font.size = SIZE_BODY
            run.italic = True
            continue

        # 普通文本 → 累积到段落
        # 清理内联markdown
        clean = line.strip()
        clean = re.sub(r'\*\*(.*?)\*\*', r'\1', clean)
        clean = re.sub(r'`(.*?)`', r'\1', clean)
        clean = re.sub(r'\[(\d+)\]', r'[\1]', clean)  # 保留引用标记

        if current_text:
            current_text.append(' ')
        current_text.append(clean)

    flush_text()
    flush_table()

def main():
    md_path = r"C:\Users\cxx\modbus-diag\论文\毕业论文_完整版.md"
    docx_path = r"C:\Users\cxx\modbus-diag\论文\毕业论文_排版版.docx"

    print("正在创建文档...")
    doc = create_document()

    print("正在添加封面...")
    add_cover_page(doc)

    print("正在解析Markdown并排版...")
    parse_markdown_and_build(doc, md_path)

    print(f"正在保存到 {docx_path} ...")
    doc.save(docx_path)

    # 统计
    size = os.path.getsize(docx_path)
    print(f"完成！文件大小: {size/1024:.1f} KB")
    print(f"文件路径: {docx_path}")

if __name__ == '__main__':
    main()
